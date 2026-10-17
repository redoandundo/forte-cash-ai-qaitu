from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class VideoRecord:
    video_id: int
    display_name: str
    path: Path
    has_violation: bool
    timecode_sec: float | None
    raw_label: str

    def as_row(self) -> dict[str, object]:
        row = asdict(self)
        row["path"] = str(self.path)
        return row


def video_id(value: object) -> int:
    match = re.search(r"(\d+)", str(value))
    if not match:
        raise ValueError(f"Could not find a numeric video id in {value!r}")
    return int(match.group(1))


def parse_timecode(value: object) -> float | None:
    if value is None or not str(value).strip():
        return None

    text = str(value).lower().replace(",", ".")
    minute_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:мин|m)", text)
    second_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:сек|с(?:\b|$)|s(?:ec)?)", text)

    minutes = float(minute_match.group(1)) if minute_match else 0.0
    if second_match:
        seconds = float(second_match.group(1))
    elif not minute_match:
        number = re.search(r"\d+(?:\.\d+)?", text)
        seconds = float(number.group()) if number else 0.0
    else:
        seconds = 0.0
    return minutes * 60 + seconds


def _is_positive(label: object) -> bool:
    text = str(label or "").strip().lower()
    return bool(text) and "нет" not in text and "без" not in text


def _find_labels(data_root: Path) -> Path:
    candidates = sorted(data_root.rglob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No .xlsx label file found under {data_root}")
    return candidates[0]


def _read_labels(path: Path) -> dict[int, tuple[str, bool, float | None, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    labels: dict[int, tuple[str, bool, float | None, str]] = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        name = str(row[0]).strip()
        raw_timecode = row[1] if len(row) > 1 else None
        raw_label = str(row[2] if len(row) > 2 else "").strip()
        labels[video_id(name)] = (
            name,
            _is_positive(raw_label),
            parse_timecode(raw_timecode),
            raw_label,
        )
    return labels


def build_index(data_root: str | Path) -> list[VideoRecord]:
    root = Path(data_root).resolve()
    labels = _read_labels(_find_labels(root))
    videos = sorted(root.rglob("*.mov"), key=lambda path: video_id(path.stem))

    records: list[VideoRecord] = []
    for path in videos:
        current_id = video_id(path.stem)
        if current_id not in labels:
            raise ValueError(f"No label found for video {path.name}")
        name, positive, timecode, raw_label = labels[current_id]
        records.append(
            VideoRecord(
                video_id=current_id,
                display_name=name,
                path=path,
                has_violation=positive,
                timecode_sec=timecode,
                raw_label=raw_label,
            )
        )

    missing = sorted(set(labels) - {record.video_id for record in records})
    if missing:
        raise ValueError(f"Videos missing for label ids: {missing}")
    return records


def rows(records: Iterable[VideoRecord]) -> list[dict[str, object]]:
    return [record.as_row() for record in records]

